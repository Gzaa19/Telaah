import io
from unittest.mock import patch, MagicMock
import pytest
from docx import Document as DocxDocument

from backend.src.documents.extractor import (
    TextExtractor,
    ScannedPDFError,
    EncryptedPDFError,
    UnsupportedFileTypeError,
)


def test_extract_docx_real():
    extractor = TextExtractor()

    # Create a real in-memory DOCX file
    doc = DocxDocument()
    doc.add_heading("Judul Riset Ilmiah", level=1)
    doc.add_paragraph("Ini adalah abstrak dari manuskrip ilmiah.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metode"
    table.cell(0, 1).text = "Akurasi"
    table.cell(1, 0).text = "RAG Baseline"
    table.cell(1, 1).text = "92.5%"

    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()

    result = extractor.extract_text(docx_bytes, "docx")
    assert result.total_pages >= 1
    assert "Judul Riset Ilmiah" in result.full_text
    assert "Ini adalah abstrak dari manuskrip ilmiah." in result.full_text
    assert "RAG Baseline" in result.full_text
    assert "92.5%" in result.full_text
    assert len(result.pages) >= 1


def test_extract_pdf_mock():
    extractor = TextExtractor()

    mock_page1 = MagicMock()
    mock_page1.extract_text.return_value = "Page 1: Introduction to AI research and literature survey.\n" * 5
    mock_page2 = MagicMock()
    mock_page2.extract_text.return_value = "Page 2: Methodology and experimental evaluation.\n" * 5

    mock_pdf = MagicMock()
    mock_pdf.pages = [mock_page1, mock_page2]
    mock_pdf.__enter__.return_value = mock_pdf

    with patch("pdfplumber.open", return_value=mock_pdf):
        result = extractor.extract_text(b"%PDF-1.4 mock bytes", "pdf")
        assert result.total_pages == 2
        assert len(result.pages) == 2
        assert result.pages[0]["page_number"] == 1
        assert "Introduction to AI research" in result.pages[0]["content"]
        assert result.pages[1]["page_number"] == 2
        assert "Methodology and experimental evaluation" in result.pages[1]["content"]


def test_extract_scanned_pdf_error():
    extractor = TextExtractor()

    mock_page = MagicMock()
    mock_page.extract_text.return_value = ""  # No text (image only)

    mock_pdf = MagicMock()
    mock_pdf.pages = [mock_page, mock_page]
    mock_pdf.__enter__.return_value = mock_pdf

    with patch("pdfplumber.open", return_value=mock_pdf):
        with pytest.raises(ScannedPDFError) as exc_info:
            extractor.extract_text(b"%PDF-1.4 scanned", "pdf")
        assert "PDF does not contain extractable text" in str(exc_info.value)


def test_unsupported_file_type_error():
    extractor = TextExtractor()
    with pytest.raises(UnsupportedFileTypeError):
        extractor.extract_text(b"some content", "png")
